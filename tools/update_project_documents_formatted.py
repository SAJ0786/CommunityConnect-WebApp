from pathlib import Path
import shutil

from docx import Document


ROOT = Path(r"C:\Users\sajja\OneDrive - SIZA Family\1. Professional\AI Apps\Community Event")
BASE = Path("output/docs")
BASE.mkdir(parents=True, exist_ok=True)

ORIGINAL_DOCX = ROOT / "Community_Events_Australia_Project_Report_2026-V90.docx"
FORMATTED_DOCX = BASE / "Community_Events_Australia_Project_Report_2026-V91-Formatted.docx"
ROADMAP_HTML_SRC = Path("roadmap_final.html")
FORMATTED_ROADMAP_HTML = BASE / "Community_Events_Roadmap_A3_Landscape_Updated_Formatted.html"


def copy_run_format(src, dst):
    dst.bold = src.bold
    dst.italic = src.italic
    dst.underline = src.underline
    dst.font.name = src.font.name
    dst.font.size = src.font.size
    if src.font.color and src.font.color.rgb:
        dst.font.color.rgb = src.font.color.rgb


def set_paragraph_text_preserve(paragraph, text):
    existing = paragraph.runs[0] if paragraph.runs else None
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.add_run(text)
    if existing:
        copy_run_format(existing, run)


def replace_exact(doc, mapping):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in mapping:
            set_paragraph_text_preserve(paragraph, mapping[text])


def replace_prefix(doc, prefix, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_paragraph_text_preserve(paragraph, text)
            return


def cell_set_text(cell, text):
    paragraph = cell.paragraphs[0]
    set_paragraph_text_preserve(paragraph, text)


def update_table_row(table, key, value, new_key=None):
    for row in table.rows:
        if row.cells[0].text.strip() == key:
            if new_key:
                cell_set_text(row.cells[0], new_key)
            cell_set_text(row.cells[1], value)
            return


def update_docx():
    shutil.copy2(ORIGINAL_DOCX, FORMATTED_DOCX)
    doc = Document(FORMATTED_DOCX)

    mapping = {
        "~11,000 lines (frontend) + ~2,400 lines (Cloud Functions)":
            "~11,000 lines frontend + Cloud Functions + native streaming relay",
        "Maintain complete privacy controls — private events visible only to registered users, poster images hidden from guests.":
            "Maintain complete privacy controls — guests can browse event cards, but private-event full addresses, host phone numbers, uploaded posters, directions, and raw Firestore event documents are protected.",
        "src/services/ — 13 service modules: eventService, hijriService, imageService, organisationsService, hmsService, streamService, userService, youtubeService, etc.":
            "src/services/ — service modules for events, Hijri conversion, images, organisations, native streaming, notifications, settings, users, and YouTube integration. Obsolete 100ms/HMS and Daily streaming services have been removed.",
        "Event management — onEventCreated (auto-creates 100ms room), recalculateHijriEvents, transferEventOwnership":
            "Event management — onEventCreated notifications, recalculateHijriEvents, transferEventOwnership, safe public event listing for guests",
        "Streaming — startYouTubeStream, stopYouTubeStream, startHLSStream, stopHLSStream, hmsWebhook":
            "Streaming — nativeStreamStart, nativeStreamCheckHealth, nativeStreamEndEvent, cleanupStaleLiveEvents, Cloud Run relay health/session handling",
        "Firebase Authentication handles all identity management. Two providers: Google OAuth (one-tap sign-in) and Phone number (OTP via SMS). Four roles: user, admin, superAdmin, guest. Guest mode allows browsing without login.":
            "Firebase Authentication now uses phone OTP as the only registered-user sign-in method. Google and Microsoft sign-in were removed, with an account migration flow that links existing email-based profiles to the new verified phone account. Four roles remain: user, admin, superAdmin, and guest. Guest mode uses anonymous authentication for safe browsing without access to restricted features.",
        "Private events restrict poster image to logged-in users — guests always see the organisation fallback logo.":
            "Private events restrict full address, directions, host phone number, and poster image to registered users. Guests see suburb-level location and organisation fallback imagery only.",
        "The Hijri calendar uses community-set monthly overrides as anchors. Each override specifies the Gregorian date of the 1st of a Hijri month. Un-anchored months cascade 30 days from the most recent prior override.":
            "The Hijri calendar uses community-set monthly overrides as anchors. Each override stores the Gregorian date of the 1st of a Hijri month, preserving correction history so later monthly adjustments do not drift into impossible dates.",
        "On override save, recalculateHijriEvents Cloud Function recalculates all Hijri-entered events.":
            "On override save, recalculateHijriEvents recalculates Hijri-entered events while preserving their Hijri anchor. Gregorian-entered events keep their Gregorian anchor and receive recalculated displayed Hijri dates.",
        "Events entered as Hijri store hijriDay/Month/Year for floating Gregorian recalculation.":
            "Events entered as Hijri store hijriDay, hijriMonth, hijriYear, and enteredAsHijri for floating Gregorian recalculation after calendar correction.",
        "Client hijriService.js and server functions/index.js implement identical conversion algorithms.":
            "Client hijriService.js and server functions/index.js use the same override-anchor model and 30-day month cascade from the latest valid correction.",
        "Admin Repair tool for fixing series date mismatches using array index as offset.":
            "Admin repair tooling supports correction of Hijri-entered events and series records after monthly calendar adjustments.",
        "Live streaming uses a multi-layer architecture: 100ms.live (WebRTC), YouTube OAuth2 (broadcast creation), and RTMP beaming (YouTube ingest).":
            "Live streaming now uses the native browser streaming path: getUserMedia and MediaRecorder in the PWA, a WebSocket connection to the Cloud Run native-stream-relay, FFmpeg normalization, and YouTube RTMPS ingest.",
        "100ms.live Integration": "Native Relay Integration",
        "Each event has a dedicated 100ms room (liveRoomId) auto-created by onEventCreated. The broadcaster joins via an iframe using a JWT generated client-side with HMAC-SHA256 via the Web Crypto API.":
            "The host opens a camera pre-check frame, selects camera/microphone and portrait or landscape mode, then starts streaming. The browser sends media chunks to the native relay; no 100ms rooms, iframe meeting URLs, or client-side streaming provider tokens are used.",
        "When Go Live is tapped, startYouTubeStream Cloud Function: (1) refreshes OAuth token; (2) creates a YouTube LiveBroadcast with the event title; (3) creates a LiveStream resource; (4) binds them; (5) starts RTMP beam from 100ms; (6) stores liveWatchUrl on the event document. The beam bot joins via /meeting/ path to avoid preview screen blocking.":
            "When Go Live is tapped, nativeStreamStart refreshes the OAuth token, creates and binds the YouTube LiveBroadcast/LiveStream resources, creates a nativeStreamSessions document, and returns the relay WebSocket URL plus the YouTube watch URL. nativeStreamCheckHealth verifies stream/broadcast state and transitions ready/testing broadcasts to live when YouTube reports active ingest.",
        "Stream Persistence & iOS Limitation": "Stream Lifecycle, Orientation, and Host Controls",
        "iOS Safari suspends camera/mic in invisible iframes. Off-screen CSS positioning (left: -99999px) keeps the iframe in the render tree while minimised. A prominent warning banner informs hosts to keep the screen open while live, as this is an iOS platform limitation.":
            "The host frame remains open after ending and shows the Stream Ended confirmation. During streaming, hosts receive an on-screen orientation warning if the phone is rotated away from the selected start orientation. The relay normalizes output to 16:9 YouTube ingest, improves audio quality, and preserves portrait content with side bars where appropriate.",
        "cleanupStaleLiveEvents runs every 10 minutes. hmsWebhook handles real-time cleanup on beam.stopped, session.close, and peer.leave events.":
            "cleanupStaleLiveEvents runs as a scheduled safety net. nativeStreamEndEvent clears event live state immediately when the host ends the stream, marks the session stopped, and attempts to complete the YouTube broadcast.",
        "Events: public read (status == active, hidden != true). Admin write requires role check. Users: owner or admin read only. Organisations: public read, admin write.":
            "Events: registered users can read full event documents. Guests no longer read raw event documents directly; they use getPublicEvents, which returns a sanitized field set. Users: owner or admin read only. Organisations: public read, admin write.",
        "6.5 100ms JWT Generation": "6.5 Obsolete Provider Cleanup",
        "100ms room tokens generated client-side using HMAC-SHA256 via the Web Crypto API. App secret stored in Firebase Remote Config / environment secrets, not the client bundle.":
            "The obsolete 100ms/HMS and Daily streaming code paths, admin backfill action, old callable functions, and hardcoded provider credentials were removed. Active streaming now uses only the native relay and YouTube API.",
        "The application is live in production at communityevents.siza.info at version 90, actively used by the community, and represents a strong foundation for continued development.":
            "The application is live in production at communityevents.siza.info at version 91, actively used by the community, and represents a strong foundation for continued development.",
    }
    replace_exact(doc, mapping)
    replace_prefix(
        doc,
        "At the time of this report the application is at version 90",
        "At the time of this report the application is at version 91, representing over 90 iterative releases spanning initial development through a complete v2 GUI redesign, dynamic organisation management, phone-only login migration, guest privacy hardening, corrected Hijri override storage, and native YouTube live streaming.",
    )
    replace_prefix(
        doc,
        "The application follows a serverless, Firebase-centric architecture",
        "The application follows a serverless, Firebase-centric architecture with a React SPA as the frontend client, Firebase Cloud Functions as the backend processing layer, Firestore as the primary database, Cloud Run as the native streaming relay, and YouTube as the broadcast destination.",
    )
    replace_prefix(
        doc,
        "A complete visual redesign was undertaken across versions 67",
        "A complete visual redesign and platform hardening cycle was undertaken across versions 67-91, transforming the original interface into a polished, image-led community app with stronger privacy, login, calendar, and streaming architecture.",
    )
    replace_prefix(
        doc,
        "100ms audio quality",
        "Streaming observability — continue monitoring native relay logs, YouTube transition health, mobile audio quality, and host orientation prompts during live community events.",
    )
    replace_prefix(
        doc,
        "Stale live cleanup speed",
        "Guest privacy verification — periodically test event cards, AI search, share messages, calendar views, and Firestore rules to confirm guests only receive sanitized event data.",
    )

    t0, t1, t2, t3, t4, t5, t6 = doc.tables
    update_table_row(t0, "Version at Report", "v91")
    update_table_row(t0, "Total Codebase", "~11,000 lines frontend + Cloud Functions + native streaming relay")
    update_table_row(t1, "Auth", "Firebase Authentication — Phone OTP for registered users; anonymous guest mode; email-based profile migration for former Google/Microsoft users")
    update_table_row(t1, "Streaming", "Native browser stream -> Cloud Run WebSocket relay -> FFmpeg -> YouTube RTMPS")
    update_table_row(t1, "Broadcasting", "YouTube Data API v3 + OAuth2 — auto-create broadcasts, bind streams, transition live/complete")
    update_table_row(t2, "events", "eventDate, eventType, hostName, audienceType, organiserType, isLive, liveWatchUrl, imageUrl, imagePath, hijriDate, hijriDay/Month/Year, enteredAsHijri, isSeries, hidden, status, createdByUserId")
    update_table_row(t3, "Firestore", "firebase deploy --only firestore:rules — registered-user raw event access; guest access via getPublicEvents")
    update_table_row(t5, "100ms.live", "Browser MediaRecorder + WebSocket relay + FFmpeg output to YouTube RTMPS", "Native streaming relay")
    update_table_row(t5, "Web Crypto API", "Host-only orientation mismatch detection for mobile streaming", "Device Orientation API")
    update_table_row(t6, "v60 (baseline)", "Text-only event cards, basic Hijri calendar, early live streaming, YouTube OAuth")
    row = t6.add_row()
    cell_set_text(row.cells[0], "v91")
    cell_set_text(row.cells[1], "Phone-only login migration; duplicate migrated profiles filtered; guest privacy hardened with sanitized Cloud Function; Hijri correction history fixed; native streaming orientation/audio improvements; obsolete 100ms/HMS/Daily code removed; GitHub source pushed.")

    doc.save(FORMATTED_DOCX)


def update_roadmap_html():
    text = ROADMAP_HTML_SRC.read_text(encoding="utf-8")
    replacements = {
        "Google Authentication": "Phone OTP Authentication",
        "Secure login with Google account for all users": "Phone login with email profile migration",
        "Secure Google login": "Phone OTP login",
        "Hijri calendar auto-calculation": "Hijri calendar stored corrections",
        "Native relay video streaming": "Native relay video streaming",
        "Per-event dedicated rooms": "Cloud Run WebSocket relay",
        "Multiple simultaneous streams": "YouTube RTMPS ingest",
        "YouTube OAuth2 auto-connection": "YouTube OAuth2 Production mode",
        "Stream persistence (return button)": "Host orientation prompt",
        "Auto-end stale events (webhook)": "Auto-end stale events",
        "Room created on event creation": "No 100ms/HMS dependency",
        "Live Streaming (Native relay)": "Live Streaming (Native relay)",
        "Per-event rooms, multiple simultaneous streams": "Browser -> relay -> FFmpeg -> YouTube",
        "Auto Cleanup": "Stream Cleanup",
        "Stale live events auto-ended via webhook": "Immediate end + scheduled safety net",
        "Privacy Controls": "Guest Privacy Controls",
        "Role-based sharing matrix for all event types": "Sanitized guest cards, AI search and share messages",
        "Native camera for streaming": "Mobile streaming permissions",
        "Add New OAuth Credentials": "Confirm Platform Settings",
        "In Google Cloud Console, create a new OAuth client ID of type Android with your app's package name and SHA-1 fingerprint from the signing key. This allows Google login to work in the native app.":
            "Confirm Firebase, push notification, camera and microphone settings for each packaged app. Phone OTP remains the app login method.",
        "Firebase Authentication": "Firebase Phone OTP",
        "Native relay</div>": "Native relay + FFmpeg</div>",
        "v52 Â· June 2026": "v91 · June 2026",
        "v52 · June 2026": "v91 · June 2026",
        "Powered by React, Firebase, Native relay, YouTube API &amp; Claude AI": "Powered by React, Firebase, Native relay, YouTube API &amp; Claude AI",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    # Keep the original visual source but add a current-status note.
    text = text.replace(
        "Current status as of June 2026",
        "Current status as of June 2026 · Updated for v91",
    )
    FORMATTED_ROADMAP_HTML.write_text(text, encoding="utf-8")


if __name__ == "__main__":
    update_docx()
    update_roadmap_html()
    print(FORMATTED_DOCX)
    print(FORMATTED_ROADMAP_HTML)
